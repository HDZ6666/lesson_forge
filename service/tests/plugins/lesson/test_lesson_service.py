"""教案生成服务在合成题库上的端到端测试（计划 4 验收 A1/B1~B5/D1/D2）。"""

from __future__ import annotations

import io
import struct
import zlib
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
import pytest_asyncio
import yaml
from docx import Document
from sqlalchemy import BigInteger, select
from sqlalchemy.dialects.mysql import LONGTEXT
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.ext.compiler import compiles

from exceptions.exception import ServiceException
from plugins.bank.entity.do.bank_do import (
    QbCatalogNode,
    QbQuestion,
    QbQuestionAsset,
    QbQuestionCatalog,
    QbTextbook,
    QbTextbookVersion,
)
from plugins.core.manifest.schema import PluginManifestFactory
from plugins.lesson.entity.do.lesson_do import LpGeneration, LpGenerationItem, LpPickRule, LpScheme
from plugins.lesson.entity.vo.lesson_vo import (
    DefaultRulesRequestModel,
    GenerationPageQueryModel,
    LessonTextModel,
    RuleInputModel,
    SchemeInputModel,
    SimulateRequestModel,
    TextUpdateModel,
)
from plugins.lesson.service.lesson_service import LessonService
from plugins.lesson.service.word_export import build_lesson_docx, html_blocks


@compiles(BigInteger, 'sqlite')
def _compile_bigint(_type: BigInteger, _compiler: object, **_kwargs: object) -> str:
    return 'INTEGER'


@compiles(LONGTEXT, 'sqlite')
def _compile_longtext(_type: LONGTEXT, _compiler: object, **_kwargs: object) -> str:
    return 'TEXT'


TABLES = [
    QbTextbookVersion.__table__,
    QbTextbook.__table__,
    QbCatalogNode.__table__,
    QbQuestion.__table__,
    QbQuestionCatalog.__table__,
    QbQuestionAsset.__table__,
    LpScheme.__table__,
    LpPickRule.__table__,
    LpGeneration.__table__,
    LpGenerationItem.__table__,
]

TEXTBOOK_ID = 1
CHAPTER_ID = 10
SECTION_ID = 11
POINT_1 = 21
POINT_2 = 22
POINT_3 = 23
QUESTIONS_PER_POINT = 24
TYPES = ['单选题', '填空题', '解答题-应用题', '解答题']
NOW = datetime(2026, 9, 7, 12, 0, 0)


def _png_bytes() -> bytes:
    """一个 2x2 的红色 PNG，用于素材嵌入测试。"""
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * 2 for _ in range(2))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (
        b'\x89PNG\r\n\x1a\n'
        + chunk(b'IHDR', struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0))
        + chunk(b'IDAT', zlib.compress(raw))
        + chunk(b'IEND', b'')
    )


def _node(catalog_id: int, parent_id: int | None, node_type: str, name: str) -> QbCatalogNode:
    return QbCatalogNode(
        catalog_id=catalog_id,
        textbook_id=TEXTBOOK_ID,
        parent_id=parent_id,
        source_provider='test',
        external_node_id=str(catalog_id),
        node_name=name,
        node_type=node_type,
        sort_order=catalog_id,
        status='0',
        create_time=NOW,
        update_time=NOW,
    )


def _question(question_id: int, index: int, point: int) -> QbQuestion:
    ease = round(0.3 + (index % 12) * 0.06, 4)  # 0.30 ~ 0.96
    return QbQuestion(
        question_id=question_id,
        source_provider='test',
        external_question_id=str(question_id),
        paper_name=f'{2020 + index % 7}年{"江苏南京" if index % 3 == 0 else "山东济南"}{"中考真题" if index % 5 == 0 else "期末"}',
        year=2020 + index % 7,
        province='江苏' if index % 3 == 0 else '山东',
        city='南京' if index % 3 == 0 else '济南',
        exam_type='中考' if index % 5 == 0 else '期末',
        question_type=TYPES[index % len(TYPES)],
        difficulty='容易' if ease >= 0.8 else '适中' if ease >= 0.55 else '困难',
        difficulty_coefficient=ease,
        content_html=(
            f'<p>第 {question_id} 题：已知 x<sup>2</sup> 与知识点 {point}。</p>'
            + ('<p><img data-asset-key="img-1" /></p>' if index % 4 == 0 else '')
            + ('<table><tr><td>a</td><td>b</td></tr><tr><td>1</td><td>2</td></tr></table>' if index % 6 == 0 else '')
        ),
        plain_text=f'第 {question_id} 题',
        answer_status='not_collected',
        subquestion_count=1 + index % 3,
        assemble_count=index * 3,
        status='0',
        create_time=NOW,
        update_time=NOW,
    )


def _relation(question_id: int, catalog_id: int, evidence: str = 'query_membership') -> QbQuestionCatalog:
    return QbQuestionCatalog(
        question_id=question_id,
        catalog_id=catalog_id,
        relation_type='query_catalog',
        evidence_type=evidence,
        sort_order=0,
        create_time=NOW,
        update_time=NOW,
    )


@pytest_asyncio.fixture
async def db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    engine = create_async_engine('sqlite+aiosqlite:///:memory:')
    async with engine.begin() as connection:
        await connection.run_sync(lambda sync: QbQuestion.metadata.create_all(sync, tables=TABLES))
    factory = async_sessionmaker(engine, expire_on_commit=False)
    asset_root = tmp_path / 'bank'
    (asset_root / 'img').mkdir(parents=True)
    (asset_root / 'img' / 'red.png').write_bytes(_png_bytes())
    monkeypatch.setenv('BANK_ASSET_ROOT', str(asset_root))
    async with factory() as session:
        session.add(
            QbTextbookVersion(
                version_id=1, source_provider='test', version_name='测试版', sort_order=0, status='0', create_time=NOW, update_time=NOW
            )
        )
        session.add(
            QbTextbook(
                textbook_id=TEXTBOOK_ID,
                version_id=1,
                source_provider='test',
                external_textbook_id='tb',
                textbook_name='九年级上册',
                sort_order=0,
                status='0',
                create_time=NOW,
                update_time=NOW,
            )
        )
        session.add_all(
            [
                _node(CHAPTER_ID, None, 'chapter', '第二章 一元二次方程'),
                _node(SECTION_ID, CHAPTER_ID, 'section', '2 用配方法求解一元二次方程'),
                _node(POINT_1, SECTION_ID, 'point', '配方法'),
                _node(POINT_2, SECTION_ID, 'point', '直接开平方法'),
                _node(POINT_3, SECTION_ID, 'point', '一元二次方程的根'),
            ]
        )
        question_id = 1000
        for point in (POINT_1, POINT_2, POINT_3):
            for index in range(QUESTIONS_PER_POINT):
                question_id += 1
                session.add(_question(question_id, index, point))
                session.add(_relation(question_id, point, 'provider_asserted' if index % 2 else 'query_membership'))
                session.add(_relation(question_id, SECTION_ID))
                if index % 8 == 7:  # 部分题目关联第二个知识点
                    session.add(_relation(question_id, POINT_3 if point != POINT_3 else POINT_1))
                if index % 4 == 0:
                    session.add(
                        QbQuestionAsset(
                            question_id=question_id,
                            source_provider='test',
                            asset_key='img-1',
                            asset_type='image',
                            storage_path='img/red.png',
                            mime_type='image/png',
                            download_status='downloaded',
                            sort_order=0,
                            create_time=NOW,
                            update_time=NOW,
                        )
                    )
        await session.commit()
        yield session
    await engine.dispose()


def _scheme(seed: int | None = 12345, **overrides: Any) -> SchemeInputModel:
    payload: dict[str, Any] = {
        'textbookId': TEXTBOOK_ID,
        'lessonType': 'new',
        'title': '配方法解一元二次方程',
        'preferredProvince': '江苏',
        'preferredCity': '南京',
        'preferredExamTypes': ['中考'],
        'seed': seed,
        'rules': [
            {'phase': 'intro', 'catalogIds': [POINT_1, POINT_2], 'questionCount': 1, 'questionTypes': ['单选题', '填空题'], 'easeMin': 0.8, 'easeMax': 1.0, 'maxExtraPoints': 1},
            {'phase': 'example', 'catalogIds': [POINT_1], 'questionCount': 2, 'easeMin': 0.55, 'easeMax': 0.85},
            {'phase': 'practice', 'catalogIds': [POINT_2], 'questionCount': 3, 'easeMin': 0.55, 'easeMax': 0.9},
            {'phase': 'extend', 'catalogIds': [POINT_1, POINT_2], 'questionCount': 1, 'questionTypes': ['解答题'], 'easeMin': 0.0, 'easeMax': 0.55},
            {'phase': 'homework', 'catalogIds': [POINT_1, POINT_2], 'questionCount': 4, 'easeMin': 0.4, 'easeMax': 1.0},
        ],
    }
    payload.update(overrides)
    return SchemeInputModel.model_validate(payload)


def _question_ids(detail: dict[str, Any]) -> list[int]:
    return [item['questionId'] for phase in detail['phases'] for item in phase['items']]


def test_manifest_is_valid_and_declares_menu_permissions() -> None:
    manifest_path = Path(__file__).resolve().parents[3] / 'plugins' / 'lesson' / 'plugin.yaml'
    manifest = PluginManifestFactory.create(yaml.safe_load(manifest_path.read_text(encoding='utf-8')))
    assert manifest.id == 'lesson'
    assert manifest.backend.module == 'plugins.lesson'
    assert 'lesson:plan:list' in manifest.permission_codes
    assert manifest.frontend.menus[0].children[0].component == 'plugin/lesson/editor/index'
    assert [plugin.id for plugin in manifest.dependencies.plugins] == ['bank']
    assert any(item.startswith('python-docx') for item in manifest.dependencies.python)


@pytest.mark.asyncio
async def test_default_rules_expand_section_to_points(db: Any) -> None:
    rules = await LessonService.default_rules(db, DefaultRulesRequestModel(lessonType='new', catalogIds=[SECTION_ID]))
    homework = next(rule for rule in rules if rule['phase'] == 'homework')
    assert homework['catalogIds'] == [POINT_1, POINT_2, POINT_3]
    assert sum(1 for rule in rules if rule['phase'] == 'example') == 3


@pytest.mark.asyncio
async def test_simulate_reports_candidate_counts_per_level(db: Any) -> None:
    rows = await LessonService.simulate(
        db,
        SimulateRequestModel(
            textbookId=TEXTBOOK_ID,
            rules=[
                RuleInputModel(phase='practice', catalogIds=[POINT_1], questionCount=3, easeMin=0.9, easeMax=1.0),
                RuleInputModel(
                    phase='practice', catalogIds=[POINT_1], questionCount=20, allowImage=False, allowTable=False
                ),
            ],
        ),
    )
    strict, loose = rows
    assert strict['candidateCounts'][0] <= strict['candidateCounts'][1] <= strict['candidateCounts'][3]
    assert strict['candidateCounts'][3] == QUESTIONS_PER_POINT + 3  # 24 题 + 3 题通过第二知识点关联
    assert strict['satisfiedLevel'] is not None
    # 含图 6 题、含表格另有 2 题被硬排除，兜底也只剩 19 题。
    assert loose['candidateCounts'][3] == 19
    assert loose['satisfiedLevel'] is None
    assert '需要 20 题' in loose['message']


@pytest.mark.asyncio
async def test_generate_is_reproducible_unique_and_respects_preferences(db: Any) -> None:
    first = await LessonService.generate(db, _scheme(), 'tester')
    second = await LessonService.generate(db, _scheme(), 'tester')
    assert _question_ids(first) == _question_ids(second)
    assert first['generationId'] != second['generationId']
    assert first['status'] == 'ready'
    assert len(_question_ids(first)) == 11
    assert len(set(_question_ids(first))) == 11
    assert [phase['phase'] for phase in first['phases']] == ['intro', 'example', 'practice', 'extend', 'homework']
    assert first['lesson']['title'] == '配方法解一元二次方程'
    assert first['lesson']['knowledgePoints'][0]['path'].startswith('第二章 一元二次方程 › ')
    assert first['lesson']['text']['objectives'].startswith('1. 理解配方法、直接开平方法')

    other_seed = await LessonService.generate(db, _scheme(seed=999), 'tester')
    assert _question_ids(other_seed) != _question_ids(first)

    intro_item = first['phases'][0]['items'][0]
    assert intro_item['question']['province'] == '江苏', '偏好省份的题应排在前面'
    assert intro_item['relaxLevel'] == 0
    for item in first['phases'][3]['items']:
        assert item['question']['questionType'].startswith('解答题')
        assert item['question']['difficultyCoefficient'] <= 0.55
    summary = {row['phase']: row for row in first['summary']}
    assert summary['homework']['needed'] == 4
    assert summary['homework']['picked'] == 4
    assert summary['homework']['short'] is False


@pytest.mark.asyncio
async def test_generate_marks_partial_when_rule_cannot_be_filled(db: Any) -> None:
    detail = await LessonService.generate(
        db,
        _scheme(
            rules=[
                {
                    'phase': 'practice',
                    'catalogIds': [POINT_1],
                    'questionCount': 20,
                    'questionTypes': ['判断题'],
                    'allowImage': False,
                    'allowTable': False,
                },
            ]
        ),
        'tester',
    )
    assert detail['status'] == 'partial'
    assert detail['summary'][0]['short'] is True
    assert detail['summary'][0]['picked'] == 19
    assert detail['summary'][0]['maxRelaxLevel'] == 3
    assert '需要 20 题' in detail['summary'][0]['message']
    assert all(item['relaxLevel'] == 3 for item in detail['phases'][0]['items']), '没有判断题时全部靠兜底选出'


@pytest.mark.asyncio
async def test_replace_item_swaps_to_unused_question_and_records_origin(db: Any) -> None:
    detail = await LessonService.generate(db, _scheme(), 'tester')
    target = detail['phases'][2]['items'][0]
    before = set(_question_ids(detail))
    updated = await LessonService.replace_item(db, detail['generationId'], target['itemId'], 'tester')
    replaced = next(item for phase in updated['phases'] for item in phase['items'] if item['itemId'] == target['itemId'])
    assert replaced['questionId'] != target['questionId']
    assert replaced['questionId'] not in before
    assert replaced['replacedFromQuestionId'] == target['questionId']
    assert len(set(_question_ids(updated))) == len(_question_ids(updated))
    with pytest.raises(ServiceException):
        await LessonService.replace_item(db, detail['generationId'], 999999, 'tester')


@pytest.mark.asyncio
async def test_regenerate_keeps_scheme_and_text_but_changes_seed(db: Any) -> None:
    detail = await LessonService.generate(db, _scheme(), 'tester')
    edited = await LessonService.update_text(
        db,
        detail['generationId'],
        TextUpdateModel(title='配方法（第 1 课时）', text=LessonTextModel(objectives='教师改写的目标')),
        'tester',
    )
    assert edited['lesson']['title'] == '配方法（第 1 课时）'
    regenerated = await LessonService.regenerate(db, detail['generationId'], 'tester')
    assert regenerated['schemeId'] == detail['schemeId']
    assert regenerated['seed'] != detail['seed']
    assert regenerated['lesson']['title'] == '配方法（第 1 课时）'
    assert regenerated['lesson']['text']['objectives'] == '教师改写的目标'
    assert regenerated['lesson']['regeneratedFrom'] == detail['generationId']
    assert len(_question_ids(regenerated)) == 11

    page = await LessonService.list_generations(db, GenerationPageQueryModel(pageNum=1, pageSize=10))
    assert page.total == 2
    assert page.rows[0]['generationId'] == regenerated['generationId']
    assert page.rows[0]['textbookName'] == '九年级上册'
    assert page.rows[0]['itemCount'] == 11

    deleted = await LessonService.delete_generations(db, [detail['generationId']])
    assert deleted == 1
    remaining = (await db.execute(select(LpGenerationItem.generation_id).distinct())).scalars().all()
    assert remaining == [regenerated['generationId']]


@pytest.mark.asyncio
async def test_generate_rejects_catalog_outside_textbook(db: Any) -> None:
    with pytest.raises(ServiceException):
        await LessonService.generate(db, _scheme(rules=[{'phase': 'practice', 'catalogIds': [999], 'questionCount': 1}]), 'tester')


@pytest.mark.asyncio
async def test_generated_lesson_keeps_question_rule_and_asset_snapshots(db: Any) -> None:
    created = await LessonService.generate(db, _scheme(), 'tester')
    image_item = next(
        item for phase in created['phases'] for item in phase['items']
        if any(asset['available'] for asset in item['question']['assets'])
    )
    question_id = image_item['questionId']
    original_html = image_item['question']['contentHtml']
    original_rules = created['rules']
    original_scheme = created['scheme']
    original_path, _ = await LessonService.asset_file(db, question_id, 'img-1', created['generationId'])

    question = await db.get(QbQuestion, question_id)
    question.content_html = '<p>正式题库已更新</p>'
    asset = (await db.execute(select(QbQuestionAsset).where(QbQuestionAsset.question_id == question_id))).scalars().first()
    asset.storage_path = 'img/new.png'
    rule = await db.get(LpPickRule, image_item['ruleId'])
    rule.rule_name = '已改动的规则'
    scheme = await db.get(LpScheme, created['schemeId'])
    scheme.scheme_name = '已改动的方案'
    await db.commit()

    history = await LessonService.detail(db, created['generationId'])
    saved_item = next(
        item for phase in history['phases'] for item in phase['items']
        if item['itemId'] == image_item['itemId']
    )
    saved_path, _ = await LessonService.asset_file(db, question_id, 'img-1', created['generationId'])
    assert saved_item['question']['contentHtml'] == original_html
    assert saved_item['question']['assets'] == image_item['question']['assets']
    assert history['rules'] == original_rules
    assert history['scheme'] == original_scheme
    assert saved_path == original_path


@pytest.mark.asyncio
async def test_word_export_contains_sections_questions_tables_and_images(db: Any) -> None:
    detail = await LessonService.generate(db, _scheme(), 'tester')
    image_items = [
        item
        for phase in detail['phases']
        for item in phase['items']
        if any(asset['available'] for asset in item['question']['assets'])
    ]
    assert image_items, '合成数据应至少选到一道带素材的题'

    async def _resolver_paths() -> dict[tuple[int, str], Path]:
        paths: dict[tuple[int, str], Path] = {}
        for item in image_items:
            for asset in item['question']['assets']:
                row = await LessonService.asset_row(db, item['questionId'], asset['assetKey'])
                paths[(item['questionId'], asset['assetKey'])] = LessonService.resolve_asset_path(row.storage_path)
        return paths

    paths = await _resolver_paths()
    content = build_lesson_docx(detail, lambda question_id, key: paths.get((question_id, key)))
    document = Document(io.BytesIO(content))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '配方法解一元二次方程' in text
    for heading in ('一、教学目标', '导入', '例题', '巩固练习', '拓展提升', '课后作业', '附录：题目来源'):
        assert heading in text
    assert '尚未采集答案' in text
    assert any('第 1' in paragraph.text and '题：已知 x' in paragraph.text for paragraph in document.paragraphs)
    assert len(document.inline_shapes) >= 1
    assert len(document.tables) >= 1  # 基本信息表


def test_html_blocks_split_paragraphs_tables_and_marks_scripts() -> None:
    blocks = html_blocks('<p>a<sup>2</sup></p><table><tr><th>x</th><th>y</th></tr><tr><td>1</td><td>2</td></tr></table><p>b</p>')
    kinds = [kind for kind, _ in blocks]
    assert kinds == ['paragraph', 'table', 'paragraph']
    assert blocks[1][1] == [['x', 'y'], ['1', '2']]
    assert blocks[0][1][1].sup is True


@pytest.mark.asyncio
async def test_historical_replacement_survives_rule_and_scheme_deletion(db: Any) -> None:
    detail = await LessonService.generate(db, _scheme(), 'tester')
    target = detail['phases'][2]['items'][0]
    rules = list((await db.execute(select(LpPickRule).where(LpPickRule.scheme_id == detail['schemeId']))).scalars())
    for rule in rules:
        await db.delete(rule)
    scheme = await db.get(LpScheme, detail['schemeId'])
    await db.delete(scheme)
    await db.commit()
    historical = await LessonService.detail(db, detail['generationId'])
    assert historical['rules'] == detail['rules']
    assert historical['scheme'] == detail['scheme']
    replaced = await LessonService.replace_item(db, detail['generationId'], target['itemId'], 'tester')
    replacement = next(item for phase in replaced['phases'] for item in phase['items'] if item['itemId'] == target['itemId'])
    assert replacement['questionId'] != target['questionId']
    assert replaced['rules'] == detail['rules']
    assert replaced['scheme'] == detail['scheme']


def test_word_table_preserves_images_superscripts_and_table_only_number(tmp_path: Path) -> None:
    image_path = tmp_path / 'formula.png'
    from PIL import Image
    Image.new('RGB', (80, 20), 'white').save(image_path)
    detail = {'lesson': {'title': '表格公式'}, 'phases': [{'phase': 'practice', 'items': [{'question': {
        'questionId': 1, 'contentHtml': '<table><tr><td>x<sup>2</sup></td><td><img data-asset-key="f1"></td></tr></table>'
    }}]}]}
    content = build_lesson_docx(detail, lambda *_: image_path)
    doc = Document(io.BytesIO(content))
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].width < 6 * 360000
    assert '题干为空' not in '\n'.join(p.text for p in doc.paragraphs)
    assert any(run.font.superscript for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs)
    with pytest.raises(ServiceException, match='无法完整导出'):
        build_lesson_docx(detail, lambda *_: None)
    image_path.write_bytes(b'not-an-image')
    with pytest.raises(ServiceException, match='图片格式'):
        build_lesson_docx(detail, lambda *_: image_path)
