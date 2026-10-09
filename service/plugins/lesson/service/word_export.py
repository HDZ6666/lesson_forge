"""教案 Word 导出：把生成详情渲染成 .docx。

题干 HTML 只保留段落、换行、上下标、表格与题图；其余标签按文本处理。
题图从 ``BANK_ASSET_ROOT`` 读取，缺失或不可用时明确拒绝不完整导出。
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from exceptions.exception import ServiceException

AssetResolver = Callable[[int, str], Path | None]
_BLOCK_TAGS = {'p', 'div', 'li', 'tr', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'section', 'article'}
_MAX_IMAGE_WIDTH_CM = 12.0
_WHITESPACE_RE = re.compile(r'[ \t\r\f\v]+')


class _Run:
    __slots__ = ('asset_key', 'kind', 'sub', 'sup', 'text')

    def __init__(self, kind: str, text: str = '', asset_key: str = '', sub: bool = False, sup: bool = False) -> None:
        self.kind = kind
        self.text = text
        self.asset_key = asset_key
        self.sub = sub
        self.sup = sup


class _HtmlBlocks(HTMLParser):
    """把题干 HTML 拆成段落块与表格块。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, Any]] = []
        self._runs: list[_Run] = []
        self._sub = 0
        self._sup = 0
        self._table: list[list[Any]] | None = None
        self._cell: list[_Run] | None = None

    def _flush(self) -> None:
        if any(run.kind == 'image' or run.text.strip() for run in self._runs):
            self.blocks.append(('paragraph', self._runs))
        self._runs = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == 'table' and self._table is None:
            self._flush()
            self._table = []
            return
        if self._table is not None:
            self._table_starttag(tag, attrs)
            return
        if tag == 'br':
            self._runs.append(_Run('text', '\n'))
        elif tag in _BLOCK_TAGS:
            self._flush()
        elif tag == 'img':
            attributes = {key: value or '' for key, value in attrs}
            self._runs.append(_Run('image', asset_key=attributes.get('data-asset-key', '').strip()))
        elif tag == 'sub':
            self._sub += 1
        elif tag == 'sup':
            self._sup += 1

    def _table_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == 'tr':
            self._table.append([])
        elif tag in {'td', 'th'}:
            if not self._table:
                self._table.append([])
            self._cell = []
        elif self._cell is not None and tag == 'img':
            attributes = dict(attrs)
            self._cell.append(_Run('image', asset_key=str(attributes.get('data-asset-key') or '').strip()))
        elif self._cell is not None and tag == 'br':
            self._cell.append(_Run('text', '\n'))
        elif tag == 'sub':
            self._sub += 1
        elif tag == 'sup':
            self._sup += 1

    def handle_endtag(self, tag: str) -> None:
        if self._table is not None:
            if tag in {'td', 'th'} and self._cell is not None:
                rich = any(run.kind == 'image' or run.sub or run.sup for run in self._cell)
                self._table[-1].append(self._cell if rich else ''.join(run.text for run in self._cell).strip())
                self._cell = None
            elif tag == 'table':
                rows = [row for row in self._table if row]
                if rows:
                    self.blocks.append(('table', rows))
                self._table = None
            elif tag == 'sub':
                self._sub = max(0, self._sub - 1)
            elif tag == 'sup':
                self._sup = max(0, self._sup - 1)
            return
        if tag in _BLOCK_TAGS:
            self._flush()
        elif tag == 'sub':
            self._sub = max(0, self._sub - 1)
        elif tag == 'sup':
            self._sup = max(0, self._sup - 1)

    def handle_data(self, data: str) -> None:
        text = _WHITESPACE_RE.sub(' ', data)
        if self._table is not None:
            if self._cell is not None:
                self._cell.append(_Run('text', text, sub=self._sub > 0, sup=self._sup > 0))
            return
        if text:
            self._runs.append(_Run('text', text, sub=self._sub > 0, sup=self._sup > 0))

    def close(self) -> None:
        super().close()
        self._flush()


def html_blocks(html: str | None) -> list[tuple[str, Any]]:
    parser = _HtmlBlocks()
    parser.feed(unescape(html) if html and '&lt;' in html and '<' not in html else (html or ''))
    parser.close()
    return parser.blocks


def _word_picture(path: Path):
    """Keep native PNG/JPEG editable as normal Word images; convert formula
    SVG and other supported raster formats to a portable embedded PNG."""
    if path.suffix.lower() == '.svg':
        from cairosvg import svg2png
        from defusedxml.ElementTree import fromstring

        source = path.read_bytes()
        root = fromstring(source)
        for element in root.iter():
            if element.tag.rsplit('}', 1)[-1] in ('script', 'foreignObject'):
                raise ValueError('不支持包含脚本或外部内容的 SVG')
            for name, value in element.attrib.items():
                if name.rsplit('}', 1)[-1] == 'href' and not value.startswith('#'):
                    raise ValueError('SVG 不允许外部引用')
                if any(
                    not ref.strip().strip('"\'').startswith('#') for ref in re.findall(r'url\((.*?)\)', value, re.I)
                ):
                    raise ValueError('SVG 不允许外部素材')
            if element.tag.rsplit('}', 1)[-1] == 'style' and re.search(r'@import|url\(', element.text or '', re.I):
                raise ValueError('SVG 样式不允许外部引用')
        return io.BytesIO(svg2png(bytestring=source, unsafe=False))
    if path.suffix.lower() in ('.webp', '.avif', '.gif', '.bmp'):
        from PIL import Image as PillowImage

        with PillowImage.open(path) as image:
            if image.width * image.height > 20_000_000:
                raise ValueError('图片像素数超过导出上限')
            result = io.BytesIO()
            image.convert('RGBA').save(result, format='PNG')
            result.seek(0)
            return result
    return str(path)


def _add_runs(
    paragraph: Any, runs: list[_Run], question_id: int, resolver: AssetResolver | None, max_width_cm: float = 6.0
) -> None:
    for run in runs:
        if run.kind == 'image':
            path = resolver(question_id, run.asset_key) if resolver and run.asset_key else None
            if path is None:
                raise ServiceException(
                    message=f'题目 {question_id} 的素材 {run.asset_key or "（无素材编号）"} 不可用，无法完整导出'
                )
            try:
                picture = _word_picture(path)
                natural_width = Image.from_file(picture).width
                if hasattr(picture, 'seek'):
                    picture.seek(0)
                paragraph.add_run().add_picture(picture, width=min(natural_width, Cm(max_width_cm)))
            except Exception as exc:
                raise ServiceException(
                    message=f'题目 {question_id} 的素材 {run.asset_key} 无法写入 Word，请检查图片格式'
                ) from exc
            continue
        docx_run = paragraph.add_run(run.text)
        if run.sub:
            docx_run.font.subscript = True
        if run.sup:
            docx_run.font.superscript = True


def _render_question(document: Any, number: int, question: dict[str, Any], resolver: AssetResolver | None) -> None:
    question_id = int(question.get('questionId') or 0)
    blocks = html_blocks(question.get('contentHtml')) or [
        ('paragraph', [_Run('text', question.get('plainText') or '（题干未采集）')])
    ]
    first = True
    for kind, payload in blocks:
        if kind == 'table':
            if first:
                document.add_paragraph().add_run(f'{number}. ').bold = True
                first = False
            columns = max(len(row) for row in payload)
            table = document.add_table(rows=len(payload), cols=columns)
            table.style = 'Table Grid'
            for row_index, row in enumerate(payload):
                for col_index, cell_text in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    if isinstance(cell_text, list):
                        _add_runs(cell.paragraphs[0], cell_text, question_id, resolver, _MAX_IMAGE_WIDTH_CM / columns)
                    else:
                        cell.text = cell_text
            continue
        paragraph = document.add_paragraph()
        if first:
            paragraph.add_run(f'{number}. ').bold = True
            first = False
        _add_runs(paragraph, payload, question_id, resolver)
    if first:
        document.add_paragraph(f'{number}. （题干为空）')
    meta = ' · '.join(
        str(value)
        for value in (
            question.get('questionType'),
            question.get('difficulty'),
            f'易度 {question.get("difficultyCoefficient")}'
            if question.get('difficultyCoefficient') is not None
            else None,
            f'{question.get("year")}年' if question.get('year') else None,
            question.get('province'),
            question.get('paperName') or question.get('sourceText'),
        )
        if value
    )
    if meta:
        note = document.add_paragraph()
        run = note.add_run(f'（{meta}）')
        run.font.size = Pt(9)
        run.italic = True


def _add_text_section(document: Any, heading: str, text: str | None) -> None:
    document.add_heading(heading, level=2)
    for line in (text or '').split('\n'):
        document.add_paragraph(line.strip() if line.strip() else '')


def build_lesson_docx(detail: dict[str, Any], resolver: AssetResolver | None = None) -> bytes:
    lesson = detail.get('lesson') or {}
    text = lesson.get('text') or {}
    document = Document()
    style = document.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)

    title = document.add_heading(str(lesson.get('title') or '教案'), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    points = lesson.get('knowledgePoints') or []
    info = document.add_table(rows=0, cols=2)
    info.style = 'Table Grid'
    for label, value in (
        ('教材', lesson.get('textbookName') or ''),
        ('课题', lesson.get('title') or ''),
        ('课型', lesson.get('lessonTypeName') or lesson.get('lessonType') or ''),
        ('课时', f'{lesson.get("durationMinutes") or 45} 分钟'),
        ('知识点', '；'.join(str(point.get('path') or point.get('name') or '') for point in points)),
    ):
        cells = info.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    _add_text_section(document, '一、教学目标', text.get('objectives'))
    _add_text_section(document, '二、教学重点', text.get('keyPoints'))
    _add_text_section(document, '三、教学难点', text.get('difficultPoints'))
    _add_text_section(document, '四、教学过程说明', text.get('process'))

    document.add_heading('五、教学环节与题目', level=1)
    number = 0
    sources: list[str] = []
    for phase in detail.get('phases') or []:
        items = phase.get('items') or []
        document.add_heading(f'{phase.get("name") or phase.get("phase")}（{len(items)} 题）', level=2)
        if not items:
            document.add_paragraph('（本环节未选到题目）')
        for item in items:
            question = item.get('question') or {'questionId': item.get('questionId')}
            number += 1
            _render_question(document, number, question, resolver)
            sources.append(
                f'{number}. '
                + ' '.join(
                    str(value)
                    for value in (
                        f'{question.get("year")}年' if question.get('year') else None,
                        question.get('province'),
                        question.get('city'),
                        question.get('examType'),
                        question.get('paperName') or question.get('sourceText') or '来源未采集',
                    )
                    if value
                )
            )

    _add_text_section(document, '六、板书要点', text.get('boardNotes'))
    _add_text_section(document, '七、作业说明', text.get('homeworkNote'))

    document.add_heading('附录：题目来源', level=2)
    for line in sources or ['（无题目）']:
        document.add_paragraph(line)
    answered = any(
        (item.get('question') or {}).get('answerHtml')
        for phase in detail.get('phases') or []
        for item in phase.get('items') or []
    )
    if not answered:
        document.add_paragraph('说明：当前题库尚未采集答案与解析，请教师在使用前补充。').runs[0].italic = True

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_question_docx(model: dict[str, Any], resolver: AssetResolver | None = None) -> bytes:
    """Question-only print model shared with the web preview."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    style = document.styles['Normal']
    style.font.name, style.font.size = '宋体', Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for name in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title = document.add_heading(model['title'], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    settings = model['printSettings']
    number = 0
    for group in model['groups']:
        if settings['showGroupTitles']:
            document.add_heading(group['name'], level=2)
        for item in group['items']:
            number += 1
            q = item['question']
            # The shared renderer adds metadata only when these fields exist.
            _render_question(
                document, number, {name: q.get(name) for name in ('questionId', 'contentHtml', 'plainText')}, resolver
            )
            for _ in range(settings['blankLines']):
                document.add_paragraph()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
